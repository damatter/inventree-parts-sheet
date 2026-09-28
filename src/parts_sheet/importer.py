"""Read workbook values as data. Never execute workbook instructions or macros."""

import csv
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from part.models import Part

ALIASES = {
    "part description": "name", "name": "name", "description": "name",
    "dicor part number": "ipn", "part number": "ipn", "ipn": "ipn",
    "# req'd": "required", "burt #": "oem_number", "dc# added to dwg": "drawing",
    "oem (usd)": "oem_usd", "local supplier (cad)": "supplier_cad",
    "dc sell (cad)": "sell_cad", "date priced": "date_priced",
    "make/model": "make_model", "material spec": "material", "size/ratio/tth": "size",
    "notes": "notes",
}


def text_cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook(content, filename):
    if len(content) > 10 * 1024 * 1024:
        raise ValidationError("The import limit is 10 MB.")
    ext = filename.lower().rsplit(".", 1)[-1]
    sheets = []
    if ext == "xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows > 10000 or sheet.ncols > 100:
                    raise ValidationError("Maximum 10,000 rows and 100 columns per sheet.")
                rows = []
                for row in sheet.get_rows():
                    rows.append([xlrd.xldate.xldate_as_datetime(c.value, book.datemode)
                                 if c.ctype == xlrd.XL_CELL_DATE else c.value for c in row])
                sheets.append((sheet.name, rows))
        finally:
            book.release_resources()
    elif ext == "xlsx":
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError("The expanded workbook is too large.")
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in book:
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                    raise ValidationError("Maximum 10,000 rows and 100 columns per sheet.")
                sheets.append((sheet.title, list(sheet.values)))
        finally:
            book.close()
    elif ext == "csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        if len(rows) > 10000 or any(len(row) > 100 for row in rows):
            raise ValidationError("Maximum 10,000 rows and 100 columns.")
        sheets = [("CSV", rows)]
    else:
        raise ValidationError("Choose an .xls, .xlsx, or UTF-8 .csv file.")
    return parse_sheets(sheets)


def parse_sheets(sheets):
    result, warnings, counters = [], [], {}
    for sheet_name, rows in sheets:
        header = None
        for index, row in enumerate(rows[:50]):
            mapping = {i: ALIASES.get(text_cell(v).casefold()) for i, v in enumerate(row)}
            if "ipn" in mapping.values() and "name" in mapping.values():
                header = (index, mapping)
                break
        if header is None:
            warnings.append(f"{sheet_name}: skipped (no part-number column).")
            continue
        for row in rows[:header[0]]:
            if any("last number used" in text_cell(v).casefold() for v in row):
                for value in row:
                    text = text_cell(value)
                    if re.fullmatch(r"100[789][0-9]{3}", text):
                        counters[text[:4]] = max(counters.get(text[:4], 0), int(text[4:]))
        for index, raw in enumerate(rows[header[0] + 1:], header[0] + 2):
            row = {key: text_cell(raw[i]) for i, key in header[1].items() if key and i < len(raw)}
            if not row.get("name") or row.get("name", "").casefold() == "part description":
                continue
            if not row.get("ipn") and not any(v for k, v in row.items() if k != "name"):
                warnings.append(f"{sheet_name} row {index}: skipped section heading {row['name']}.")
                continue
            extra = [text_cell(v) for i, v in enumerate(raw) if not header[1].get(i) and text_cell(v)]
            if extra:
                row["notes"] = " | ".join(filter(None, [row.get("notes", ""), *extra]))
            result.append({"source": f"{sheet_name}:{index}", "ipn": row.pop("ipn", ""),
                           "name": row.pop("name"), "cells": row})
    if not result:
        raise ValidationError("No part rows found. Include Part Description and DiCor Part Number columns.")
    if len(result) > 2000:
        raise ValidationError("Import at most 2,000 parts at once.")
    return result, warnings, counters


def classify(rows):
    counts = Counter(row["ipn"].casefold() for row in rows if row["ipn"])
    result = []
    for row in rows:
        matches = list(Part.objects.filter(IPN__iexact=row["ipn"])[:2]) if row["ipn"] else []
        conflict = counts[row["ipn"].casefold()] > 1 if row["ipn"] else False
        status = "conflict" if conflict or len(matches) > 1 else "matched" if matches else "new" if row["ipn"] else "unnumbered"
        result.append({**row, "status": status, "part_id": matches[0].pk if len(matches) == 1 else None,
                       "existing_name": matches[0].name if len(matches) == 1 else ""})
    return result


def exact_price(text):
    if not text:
        return None
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None
    return value if value.is_finite() and value >= 0 else None
